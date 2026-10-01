"""Builds the Junior (Easy) and Tupai (Normal) maths question banks.

Run:  python build_banks_easy_normal.py
Writes, for each bank, questions-<level>.json, a copy of it in
public/content/ for the game to load, and questions-<level>.xlsx for review.
The Tupai Hero (Hard) bank is build_bank.py's, and is not touched here.

Scope agreed with Johan, 1 Oct 2026:
  Junior (Easy): add and take away up to 3 digits, times tables to 5
    (one factor 2-5, the other 1-9). No division.
    tier 1  + and - within 20; x tables 1, 2 and 5 with factors up to 5
    tier 2  + and - within 100, some carrying and borrowing; x 2-5 tables
    tier 3  + and - with 3-digit numbers, friendly ones and some carrying;
            the harder facts up to 5 x 9
  Tupai (Normal): add and take away from 3 up to 6 digits, times tables to
    12. No division. The 7-second timer stays, so big numbers are friendly
    for mental maths (trailing zeros, round parts: 340,000 + 125,000).
    tier 1  3-digit + and -; x tables to 6
    tier 2  4- and 5-digit friendly + and -; x tables to 10
    tier 3  5- and 6-digit friendly + and -; x tables to 12, 11s and 12s
            often, plus the 7/8/9 facts kids find hardest
  - whole, positive answers only; subtraction has the bigger number first
  - "pad" marks questions short enough to type in time: an answer of at
    most 3 digits (Easy) or 4 digits (Normal)
  - wrong options are typical slips, never giveaways (see wrong_answers)
"""
import json
import random
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
# Fixed so each bank is identical on every rebuild.
SEEDS = {"easy": 20261001, "normal": 20261002}
PAD_DIGITS = {"easy": 3, "normal": 4}
SIGN = {"+": "+", "-": "−", "x": "×"}  # +, minus sign, times sign
OP_NAME = {"+": "add", "-": "sub", "x": "mul"}
OP_TYPE = {"+": "Add", "-": "Subtract", "x": "Multiply"}
TARGET_MIX = {
    "easy": {"+": 0.40, "-": 0.35, "x": 0.25},
    "normal": {"+": 0.35, "-": 0.35, "x": 0.30},
}


# ---------------------------------------------------------------- digit tools

def digits(n):
    return len(str(n))


def trailing_zeros(n):
    count = 0
    while n and n % 10 == 0:
        n //= 10
        count += 1
    return count


def carry_places(x, y):
    """Places (0 = ones) that carry into the next place when adding."""
    places, carry, place = [], 0, 0
    while x or y or carry:
        carry = 1 if x % 10 + y % 10 + carry >= 10 else 0
        if carry:
            places.append(place)
        x, y, place = x // 10, y // 10, place + 1
    return places


def borrow_places(x, y):
    """Places that borrow from the next place when taking y from x (x >= y)."""
    places, borrow, place = [], 0, 0
    while x or y:
        borrow = 1 if x % 10 - y % 10 - borrow < 0 else 0
        if borrow:
            places.append(place)
        x, y, place = x // 10, y // 10, place + 1
    return places


def smaller_from_larger(x, y):
    """The classic take-away slip: in each place, the smaller digit from the
    larger, whichever number it is in (52 - 18 -> 46)."""
    result, place = 0, 0
    while x or y:
        result += abs(x % 10 - y % 10) * 10 ** place
        x, y, place = x // 10, y // 10, place + 1
    return result


# ----------------------------------------------------------- wrong options

def addsub_candidates(x, y, op):
    """Typical slips for x + y or x - y, as {value: weight}.

    Works on the numbers with their shared trailing zeros taken off
    (340,000 + 125,000 is worked as 340 + 125), so every slip is in a place
    the child actually had to work: never 465,001 for 465,000.
    """
    unit = 10 ** min(trailing_zeros(x), trailing_zeros(y))
    x, y = x // unit, y // unit
    s = x + y if op == "+" else x - y
    n = digits(s)
    five = x % 10 == 5 or y % 10 == 5
    cands = defaultdict(int)

    def put(value, weight):
        if value > 0 and value != s:
            cands[value * unit] += weight

    for p in range(n):
        if n == 1:
            weight = 3
        elif p == 0:
            weight = 1 if five else 2  # 465 -> 466 is weak when 5s are in play
        elif p == n - 1:
            weight = 1  # first digit: a slip, but the easiest to rule out
        else:
            weight = 3  # middle digit: what stops first/last digit matching
        put(s + 10 ** p, weight)
        put(s - 10 ** p, weight)
        if p <= n - 2 or n == 1:
            put(s + 2 * 10 ** p, 1)
            put(s - 2 * 10 ** p, 1)
    if n == 1:
        put(s + 3, 1)
        put(s - 3, 1)
    if five:  # 340 + 125: dropping or doubling the 5 -> 460, 470
        put(s + 5, 2)
        put(s - 5, 2)
    if op == "+":
        for place in carry_places(x, y):
            put(s - 10 ** (place + 1), 4)  # forgot to carry
    else:
        for place in borrow_places(x, y):
            put(s + 10 ** (place + 1), 4)  # borrowed but did not pay it back
        put(smaller_from_larger(x, y), 4)
    return dict(cands)


def small_candidates(answer):
    """Tier-1 sums within 20: near misses either side (7 + 8 -> 13, 14, 16)."""
    return {answer + d: w for d, w in ((1, 3), (-1, 3), (2, 2), (-2, 2), (3, 1), (-3, 1))}


def mult_candidates(a, b, top):
    """Neighbouring table facts (7 x 8 -> 48, 63, 54), 10 out (46, 66), and
    near misses."""
    s = a * b
    cands = defaultdict(int)

    def put(value, weight):
        if value > 0 and value != s:
            cands[value] += weight

    for i, j in ((a, b + 1), (a, b - 1), (a + 1, b), (a - 1, b)):
        if min(i, j) >= 2 or min(a, b) == 1:  # 2 x 7 -> 7 (1 x 7) is no slip
            put(i * j, 4)  # one factor out by one: plus or minus a factor
    if s >= 10:  # 4 -> 14 is no slip
        put(s + 10, 3)
        put(s - 10, 3)
    for d in (1, -1, 2, -2):
        put(s + d, 1)
    reach = max(4, s // 8)
    for i in range(2, top + 1):
        for j in range(i, top + 1):
            if abs(i * j - s) <= reach:
                put(i * j, 2)  # another fact that "looks right" (56 vs 54)
    return dict(cands)


def stripped(values):
    """The values with their shared trailing zeros removed."""
    z = min(trailing_zeros(v) for v in values)
    return [v // 10 ** z for v in values]


def quality_ok(answer, wrong, level, small):
    """Does the set of four options avoid giving the answer away?

    level 0 is the strictest; each level drops one rule (see LEVEL_RULES).
    """
    if level >= 4:
        return True
    st = stripped([answer, *wrong])
    # Never the only odd one or the only even one.
    if not any(w % 2 == st[0] % 2 for w in st[1:]):
        return False
    if small or level >= 3:
        return True
    texts = [str(v) for v in st]
    if len({len(t) for t in texts}) > 1 or len(texts[0]) == 1:
        return True  # single digits, or 9 vs 10: parity is all there is
    a, ws = texts[0], texts[1:]
    if not any(w[-1] == a[-1] for w in ws):
        return False  # the last digit alone would give it away
    if not any(w[0] == a[0] for w in ws):
        return False  # the first digit alone would give it away
    if len(a) >= 3 and not any(w[0] == a[0] and w[-1] == a[-1] for w in ws):
        return False  # first + last digit would: need a middle-digit slip
    if level >= 2:
        return True
    if all(w[-1] == a[-1] for w in ws):
        return False  # mix: not every option shares the last digit
    if level >= 1:
        return True
    if any(2 * answer - w in wrong for w in wrong):
        return False  # no evenly spaced pair either side of the answer
    return True


LEVEL_RULES = [
    "all rules",
    "allowed an evenly spaced pair round the answer",
    "allowed every option to share the last digit",
    "parity only",
    "no rules",
]


def wrong_answers(answer, cands, rng, small=False):
    """Three wrong options, or fails loudly.

    Keeps the plausible candidates (same number of digits as the answer from
    20 up; within 10 of it below 20), then looks at every set of three. Of
    the sets that pass the strictest level of quality_ok that any set can
    pass, it chooses where the answer sits among the four (lowest, second,
    third, highest) at random, so position gives nothing away, and then
    prefers the most typical slips, with a little chance so the same slip
    does not always appear.
    """
    pool = {v: w for v, w in cands.items() if v > 0 and v != answer}
    if answer >= 20:
        pool = {v: w for v, w in pool.items() if digits(v) == digits(answer)}
    else:
        pool = {v: w for v, w in pool.items() if abs(v - answer) <= 10}
        if answer >= 10:  # 18 against 8 is no contest: keep 2 digits if we can
            two = {v: w for v, w in pool.items() if v >= 10}
            if len(two) >= 5:
                pool = two
    values = sorted(pool, key=lambda v: (-pool[v], abs(v - answer), v))[:18]
    sets = list(combinations(sorted(values), 3))
    for level in range(len(LEVEL_RULES)):
        by_rank = defaultdict(list)
        for s in sets:
            if quality_ok(answer, s, level, small):
                by_rank[sum(v < answer for v in s)].append(s)
        if by_rank:
            ranks = sorted(by_rank)
            rank = ranks[int(rng.random() * len(ranks))]
            best = max(by_rank[rank],
                       key=lambda s: sum(pool[v] for v in s) * (0.6 + rng.random()))
            return sorted(best), level
    raise ValueError(f"no three wrong options for {answer}: {pool}")


# ------------------------------------------------------- question families

def r(rng, lo, hi):
    return rng.randint(lo, hi)


def either(rng, x, y):
    return (x, y) if rng.random() < 0.5 else (y, x)


def has_zero(n):
    return "0" in str(n)


# (tier, op, family, how many, generator(rng) -> (x, y), accept(x, y))
# Subtraction always has the bigger number first; accept() checks the answer.
def easy_families():
    return [
        # Tier 1: within 20.
        (1, "+", "add-within-10", 30,
         lambda g: (r(g, 1, 9), r(g, 1, 9)),
         lambda x, y: 3 <= x + y <= 10),
        (1, "+", "add-within-20", 45,
         lambda g: (r(g, 1, 19), r(g, 1, 19)),
         lambda x, y: 11 <= x + y <= 20),
        (1, "-", "sub-within-10", 25,
         lambda g: (r(g, 3, 10), r(g, 1, 8)),
         lambda x, y: x - y >= 2),
        (1, "-", "sub-within-20", 40,
         lambda g: (r(g, 11, 20), r(g, 1, 18)),
         lambda x, y: x - y >= 2),
        # Tier 2: within 100.
        (2, "+", "add-2digit", 35,
         lambda g: (r(g, 10, 89), r(g, 10, 89)),
         lambda x, y: 21 <= x + y <= 99 and not carry_places(x, y)
         and not (x % 10 == 0 and y % 10 == 0) and 10 not in (x, y)),
        (2, "+", "add-2digit-carry", 40,
         lambda g: either(g, r(g, 11, 89), r(g, 2, 39)),
         lambda x, y: 21 <= x + y <= 99 and len(carry_places(x, y)) == 1),
        (2, "-", "sub-2digit", 30,
         lambda g: (r(g, 21, 99), r(g, 10, 89)),
         lambda x, y: x - y >= 10 and not borrow_places(x, y)
         and not (x % 10 == 0 and y % 10 == 0) and y != 10),
        (2, "-", "sub-2digit-borrow", 35,
         lambda g: (r(g, 21, 99), r(g, 3, 39)),
         lambda x, y: x - y >= 10 and len(borrow_places(x, y)) == 1),
        # Tier 3: 3-digit, friendly or with one carry/borrow.
        (3, "+", "add-3digit-friendly", 30, friendly_pair,
         lambda x, y: 200 <= x + y <= 999 and max(x, y) >= 100
         and not (x % 100 == 0 and y % 100 == 0) and len(carry_places(x, y)) <= 1),
        (3, "+", "add-3digit", 20,
         lambda g: (r(g, 100, 899), r(g, 100, 899)),
         lambda x, y: x + y <= 999 and not carry_places(x, y)
         and (has_zero(x) or has_zero(y))),
        (3, "+", "add-3digit-carry", 25,  # 3-digit + 2-digit: 424 + 93
         lambda g: either(g, r(g, 100, 899), r(g, 11, 99)),
         lambda x, y: x + y <= 999 and len(carry_places(x, y)) == 1),
        (3, "-", "sub-3digit-friendly", 25, friendly_take,
         lambda x, y: x - y >= 50 and not (x % 100 == 0 and y % 100 == 0)
         and len(borrow_places(x, y)) <= (2 if x % 100 == 0 else 1)),
        (3, "-", "sub-3digit", 20,
         lambda g: (r(g, 200, 999), r(g, 100, 899)),
         lambda x, y: x - y >= 50 and not borrow_places(x, y)
         and (has_zero(x) or has_zero(y))),
        (3, "-", "sub-3digit-borrow", 20,  # 3-digit - 2-digit: 452 - 28
         lambda g: (r(g, 200, 999), r(g, 11, 99)),
         # no borrowing across a zero (407 - 18) at this level
         lambda x, y: x - y >= 50 and len(borrow_places(x, y)) == 1
         and (x // 10) % 10 != 0),
    ]


def friendly_pair(g):
    """Two friendly numbers: both tens (340 + 220) or both 25s (250 + 125)."""
    if g.random() < 0.5:
        return either(g, r(g, 10, 79) * 10, r(g, 10, 59) * 10)
    return either(g, r(g, 4, 31) * 25, r(g, 2, 23) * 25)


def friendly_take(g):
    """A round number take a friendly one: 500 - 250, 750 - 125, 640 - 320."""
    if g.random() < 0.5:
        x = r(g, 4, 19) * 50
        return x, r(g, 2, x // 25 - 2) * 25
    x = r(g, 20, 99) * 10
    return x, r(g, 10, x // 10 - 5) * 10


def normal_families():
    big_carry = lambda x, y: len(carry_places(x, y)) <= 1  # noqa: E731
    big_borrow = lambda x, y: len(borrow_places(x, y)) <= 1  # noqa: E731
    sig = lambda x, y: (x - y) // 10 ** min(trailing_zeros(x), trailing_zeros(y))  # noqa: E731
    return [
        # Tier 1: 3-digit, a step up from Junior's top tier.
        (1, "+", "add-3digit-friendly", 25, friendly_pair,
         lambda x, y: 200 <= x + y <= 999 and max(x, y) >= 100
         and not (x % 100 == 0 and y % 100 == 0) and len(carry_places(x, y)) <= 1),
        (1, "+", "add-3digit", 20,
         lambda g: (r(g, 100, 899), r(g, 100, 899)),
         lambda x, y: x + y <= 999 and not carry_places(x, y)),
        (1, "+", "add-3digit-carry", 25,  # 238 + 145: the smaller one is
         lambda g: either(g, r(g, 100, 799), r(g, 100, 399)),  # a 5 or has a 0
         lambda x, y: x + y <= 999 and len(carry_places(x, y)) == 1
         and (min(x, y) % 5 == 0 or has_zero(min(x, y)))),
        (1, "-", "sub-3digit-friendly", 25, friendly_take,
         lambda x, y: x - y >= 50 and not (x % 100 == 0 and y % 100 == 0)
         and len(borrow_places(x, y)) <= (2 if x % 100 == 0 else 1)),
        (1, "-", "sub-3digit", 20,
         lambda g: (r(g, 200, 999), r(g, 100, 899)),
         lambda x, y: x - y >= 50 and not borrow_places(x, y)),
        (1, "-", "sub-3digit-borrow", 25,  # 452 - 125, 508 - 160
         lambda g: (r(g, 200, 999), r(g, 100, 399)),
         lambda x, y: x - y >= 50 and len(borrow_places(x, y)) == 1
         and (y % 5 == 0 or has_zero(y))),
        # Tier 2: 4 and 5 digits, friendly.
        (2, "+", "add-4digit-hundreds", 20,  # 4,500 + 2,300
         lambda g: (r(g, 10, 89) * 100, r(g, 10, 89) * 100),
         lambda x, y: x + y <= 9900 and big_carry(x, y)
         and not (x % 1000 == 0 and y % 1000 == 0)),
        (2, "+", "add-4digit-fifties", 15,  # 2,400 + 1,350
         lambda g: either(g, r(g, 10, 79) * 100, r(g, 10, 79) * 100 + 50),
         lambda x, y: x + y <= 9950 and big_carry(x, y)),
        (2, "+", "add-5digit-thousands", 20,  # 45,000 + 12,000
         lambda g: (r(g, 10, 89) * 1000, r(g, 10, 89) * 1000),
         lambda x, y: x + y <= 99000 and big_carry(x, y)
         and not (x % 10000 == 0 and y % 10000 == 0)),
        (2, "+", "add-5digit-halves", 15,  # 34,000 + 12,500
         lambda g: either(g, r(g, 10, 79) * 1000, r(g, 10, 79) * 1000 + 500),
         lambda x, y: x + y <= 99500 and big_carry(x, y)),
        (2, "-", "sub-4digit-hundreds", 20,  # 7,800 - 3,500
         lambda g: (r(g, 20, 99) * 100, r(g, 10, 89) * 100),
         lambda x, y: x - y >= 1000 and big_borrow(x, y) and sig(x, y) >= 10
         and not (x % 1000 == 0 and y % 1000 == 0)),
        (2, "-", "sub-4digit-fifties", 15,  # 5,000 - 1,250, 6,400 - 2,150
         lambda g: (r(g, 20, 99) * 100, r(g, 10, 79) * 100 + 50),
         lambda x, y: x - y >= 1000
         and len(borrow_places(x, y)) <= (2 if x % 1000 == 0 else 1)),
        (2, "-", "sub-5digit-thousands", 20,  # 45,000 - 12,000
         lambda g: (r(g, 20, 99) * 1000, r(g, 10, 89) * 1000),
         lambda x, y: x - y >= 1000 and big_borrow(x, y) and sig(x, y) >= 10
         and not (x % 10000 == 0 and y % 10000 == 0)),
        (2, "-", "sub-5digit-halves", 15,  # 34,000 - 12,500
         lambda g: (r(g, 20, 99) * 1000, r(g, 10, 79) * 1000 + 500),
         lambda x, y: x - y >= 1000
         and len(borrow_places(x, y)) <= (2 if x % 10000 == 0 else 1)),
        # Tier 3: 5 and 6 digits, friendly.
        (3, "+", "add-6digit-thousands", 25,  # 340,000 + 125,000; 340,000 + 85,000
         lambda g: either(g, r(g, 10, 79) * 10000,
                          r(g, 2, 129) * 5000 if g.random() < 0.7 else r(g, 11, 99) * 1000),
         lambda x, y: 100000 <= x + y < 1000000 and big_carry(x, y)
         and x % 10000 + y % 10000 != 0),
        (3, "+", "add-6digit-tenthousands", 20,  # 480,000 + 350,000
         lambda g: (r(g, 10, 89) * 10000, r(g, 10, 89) * 10000),
         lambda x, y: x + y < 1000000 and big_carry(x, y)
         and not (x % 100000 == 0 and y % 100000 == 0)),
        (3, "+", "add-5digit-hundreds", 25,  # 34,500 + 21,000
         lambda g: either(g, r(g, 100, 799) * 100, r(g, 10, 79) * 1000 + 500 * r(g, 0, 1)),
         lambda x, y: 10000 <= x + y <= 99900 and big_carry(x, y)
         and (x % 1000 != 0 or y % 1000 != 0)),
        (3, "-", "sub-6digit-round", 20,  # 800,000 - 365,000
         lambda g: (r(g, 4, 19) * 50000, r(g, 10, 179) * 5000),
         lambda x, y: x - y >= 50000 and y % 10000 != 0),
        (3, "-", "sub-6digit-thousands", 20,  # 645,000 - 320,000
         lambda g: (r(g, 40, 199) * 5000, r(g, 10, 89) * 10000),
         lambda x, y: x - y >= 50000 and x % 10000 != 0 and big_borrow(x, y)),
        (3, "-", "sub-6digit-tenthousands", 15,  # 830,000 - 450,000
         lambda g: (r(g, 20, 99) * 10000, r(g, 10, 89) * 10000),
         lambda x, y: x - y >= 50000 and big_borrow(x, y) and sig(x, y) >= 10
         and not (x % 100000 == 0 and y % 100000 == 0)),
        (3, "-", "sub-5digit-hundreds", 15,  # 56,500 - 23,000
         lambda g: (r(g, 300, 999) * 100, r(g, 10, 79) * 1000 + 500 * r(g, 0, 1)),
         lambda x, y: x - y >= 10000 and big_borrow(x, y)
         and (x % 1000 != 0 or y % 1000 != 0)),
    ]


def easy_times():
    """(a, b, tier) for each 'a x b' text. One factor 2-5 (1 in tier 1),
    the other 1-9. There are only 56 such texts, so both orders are used and
    the harder facts are split by order between tiers 2 and 3."""
    out = []
    for a in range(1, 6):
        for b in range(a, 10):
            if (a, b) == (1, 1) or (a == 1 and b > 5):
                continue
            if b <= 5 and (a in (1, 2) or b == 5):
                tiers = (1, 1)  # tables 1, 2 and 5, factors up to 5
            elif a >= 3 and b >= 6 and (a, b) != (3, 6) and a != 5:
                tiers = (3, 3)  # 3 x 7-9, 4 x 6-9
            elif a == 5:
                tiers = (2, 3)  # 5 x 6 in tier 2, 6 x 5 in tier 3
            else:
                tiers = (2, 2)  # 2 x 6-9, 3 x 3, 3 x 4, 4 x 4, 3 x 6
            out.append((a, b, tiers[0]))
            if a != b:
                out.append((b, a, tiers[1]))
    return out


def normal_times():
    """(a, b, tier) for each 'a x b' text, tables 2-12."""
    out = []
    for a in range(2, 13):
        for b in range(a, 13):
            if (a, b) == (10, 10):
                continue  # 100: no plausible same-length options below it
            if b >= 11:
                tiers = (3, 3)  # 11s and 12s
            elif b == 10:
                tiers = (2, 2)
            elif a >= 7:
                tiers = (2, 3)  # 7-9 x 7-9: one order in tier 2, the other in 3
            elif a * b <= 24:
                tiers = (1, 1)
            else:
                tiers = (1, 2)  # 6 x 7 in tier 1, 7 x 6 in tier 2
            out.append((a, b, tiers[0]))
            if a != b:
                out.append((b, a, tiers[1]))
    return out


def times_family(level, a, b):
    """The table a question belongs to: the friendliest one it is in for
    Junior (5 x 3 is a 5s fact), the 10s, 11s and 12s for Tupai."""
    for table in ((1, 2, 5) if level == "easy" else (12, 11, 10)):
        if table in (a, b):
            return f"x{table}"
    return f"x{min(a, b)}"


# ------------------------------------------------------------------- build

def fmt(n):
    return f"{n:,}"


def build(level):
    rng = random.Random(SEEDS[level])
    families = easy_families() if level == "easy" else normal_families()
    questions, texts, short = [], set(), []

    for tier, op, family, count, gen, accept in families:
        made, tries = 0, 0
        while made < count and tries < 50000:
            tries += 1
            x, y = gen(rng)
            if op == "-" and x <= y:
                continue
            if not accept(x, y):
                continue
            text = f"{fmt(x)} {SIGN[op]} {fmt(y)}"
            if text in texts:
                continue
            texts.add(text)
            answer = x + y if op == "+" else x - y
            small = level == "easy" and tier == 1
            cands = small_candidates(answer) if small else addsub_candidates(x, y, op)
            wrong, rule = wrong_answers(answer, cands, rng, small)
            questions.append({"op": op, "family": family, "tier": tier, "text": text,
                              "answer": answer, "wrong": wrong, "rule": rule})
            made += 1
        if made < count:
            short.append(f"{level} tier {tier} {family}: {made} of {count}")

    times = easy_times() if level == "easy" else normal_times()
    top = 9 if level == "easy" else 12
    for a, b, tier in times:
        text = f"{a} {SIGN['x']} {b}"
        assert text not in texts, text
        texts.add(text)
        answer = a * b
        wrong, rule = wrong_answers(answer, mult_candidates(a, b, top), rng)
        questions.append({"op": "x", "family": times_family(level, a, b), "tier": tier,
                          "text": text, "answer": answer, "wrong": wrong, "rule": rule})

    order = {"+": 0, "-": 1, "x": 2}
    questions.sort(key=lambda q: (q["tier"], order[q["op"]], q["family"], q["answer"], q["text"]))
    counters = Counter()
    prefix = level[0]
    for q in questions:
        key = (q["tier"], q["op"])
        counters[key] += 1
        q["id"] = f"{prefix}-t{q['tier']}-{OP_NAME[q['op']]}-{counters[key]:04d}"
        q["pad"] = digits(q["answer"]) <= PAD_DIGITS[level]
    return questions, short


# ------------------------------------------------------------------- check

def check(level, questions):
    """Fails loudly if anything in the bank is wrong; returns warnings for
    things that are only 'roughly right'."""
    ids, texts = set(), set()
    prefix = level[0] + "-"
    for q in questions:
        left, symbol, right = q["text"].split(" ")
        for part in (left, right):
            assert part.replace(",", "").isdigit(), q
            if int(part.replace(",", "")) >= 1000:
                assert part == fmt(int(part.replace(",", ""))), q  # separators
            else:
                assert "," not in part, q
        left, right = int(left.replace(",", "")), int(right.replace(",", ""))
        assert symbol == SIGN[q["op"]], q
        if q["op"] == "+":
            assert left + right == q["answer"], q
        elif q["op"] == "-":
            assert left > right and left - right == q["answer"], q
        else:
            assert left * right == q["answer"], q
            if level == "easy":
                assert min(left, right) <= 5 and max(left, right) <= 9, q
            else:
                assert max(left, right) <= 12, q
        assert q["answer"] > 0, q
        assert len(q["wrong"]) == 3 and len(set(q["wrong"])) == 3, q
        assert q["answer"] not in q["wrong"] and all(w > 0 for w in q["wrong"]), q
        assert all(isinstance(w, int) for w in q["wrong"]), q
        assert q["id"].startswith(prefix) and q["id"] not in ids, q
        ids.add(q["id"])
        assert q["text"] not in texts, q
        texts.add(q["text"])
        assert q["tier"] in (1, 2, 3), q
        assert q["op"] in ("+", "-", "x"), q
        assert q["pad"] == (digits(q["answer"]) <= PAD_DIGITS[level]), q

    warnings = []
    for tier in (1, 2, 3):
        rows = [q for q in questions if q["tier"] == tier]
        pad = sum(q["pad"] for q in rows)
        assert pad >= 40, f"{level} tier {tier}: only {pad} typed-eligible"
        if not 150 <= len(rows) <= 250:
            warnings.append(f"tier {tier}: {len(rows)} questions (aimed for 150-250)")
        mix = Counter(q["op"] for q in rows)
        for op, share in TARGET_MIX[level].items():
            got = mix[op] / len(rows)
            assert mix[op] > 0, f"{level} tier {tier}: no {op}"
            if abs(got - share) > 0.07:
                warnings.append(f"tier {tier} {OP_TYPE[op]}: {got:.0%} of the questions "
                                f"(aimed for {share:.0%})")
        loose = Counter(LEVEL_RULES[q["rule"]] for q in rows if q["rule"] > 0)
        for rule, n in loose.items():
            warnings.append(f"tier {tier}: {n} question(s) whose options {rule}")
    return warnings


# ------------------------------------------------------------------ output

FIELDS = ("id", "tier", "op", "family", "text", "answer", "wrong", "pad")


def write_json(level, questions):
    rows = [{k: q[k] for k in FIELDS} for q in questions]
    text = json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
    name = f"questions-{level}.json"
    (HERE / name).write_text(text, encoding="utf-8")
    # The copy the game loads.
    (HERE.parent / "public" / "content" / name).write_text(text, encoding="utf-8")


def write_xlsx(level, questions):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    header = ["ID", "Tier", "Type", "Family", "Question", "Answer",
              "Wrong 1", "Wrong 2", "Wrong 3", "Typed (number pad)"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for q in questions:
        ws.append([q["id"], q["tier"], OP_TYPE[q["op"]], q["family"], q["text"],
                   q["answer"], *q["wrong"], "yes" if q["pad"] else "no"])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=9):
        for cell in row:
            cell.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate([15, 6, 10, 24, 22, 10, 10, 10, 10, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    summary = wb.create_sheet("Summary")
    summary.append(["Tier", "Type", "Family", "Questions", "Of which typed"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for key, count, pad in summarise(questions):
        summary.append([*key, count, pad])
    for i, width in enumerate([6, 10, 24, 11, 15], 1):
        summary.column_dimensions[get_column_letter(i)].width = width
    wb.save(HERE / f"questions-{level}.xlsx")


def summarise(questions):
    groups = {}
    for q in questions:
        key = (q["tier"], OP_TYPE[q["op"]], q["family"])
        count, pad = groups.get(key, (0, 0))
        groups[key] = (count + 1, pad + q["pad"])
    return [(key, *groups[key]) for key in sorted(groups)]


def report(level, questions, warnings, short):
    title = {"easy": "Junior (Easy)", "normal": "Tupai (Normal)"}[level]
    print(f"\n{title}: {len(questions)} questions, all checked")
    print(f"  {'tier':<5}{'count':>6}  {'+':>9} {'-':>9} {'x':>9}  {'typed':>6}  "
          f"{'answers':<20} answer rank (low..high)")
    for tier in (1, 2, 3):
        rows = [q for q in questions if q["tier"] == tier]
        mix = Counter(q["op"] for q in rows)
        split = [f"{mix[op]:>3} ({mix[op] / len(rows):.0%})" for op in ("+", "-", "x")]
        ranks = Counter(sum(w < q["answer"] for w in q["wrong"]) for q in rows)
        lo, hi = min(q["answer"] for q in rows), max(q["answer"] for q in rows)
        print(f"  {tier:<5}{len(rows):>6}  {split[0]:>9} {split[1]:>9} {split[2]:>9}  "
              f"{sum(q['pad'] for q in rows):>6}  {fmt(lo) + ' - ' + fmt(hi):<20} "
              f"{[ranks[i] for i in range(4)]}")
    for line in short:
        print(f"  SHORT: {line}")
    for line in warnings:
        print(f"  note: {line}")


if __name__ == "__main__":
    for level in ("easy", "normal"):
        bank, short = build(level)
        warnings = check(level, bank)
        write_json(level, bank)
        write_xlsx(level, bank)
        report(level, bank, warnings, short)
